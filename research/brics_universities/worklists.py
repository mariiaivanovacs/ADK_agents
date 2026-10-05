"""Worklists for the local browser session: rankings (CONFIRMED) and contacts (NOT_CONFIRMED).

  python worklists.py build OUTDIR
  python worklists.py import FILLED.xlsx      # either workbook; type detected by sheet name
"""
import json
import os
import sys

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(__file__))
from manual_tasks import DATA, FILES, FILL, sheet, load  # noqa: E402

KEYS = json.load(open(os.path.join(DATA, "not_confirmed_keys.json"), encoding="utf-8"))
key = lambda r: r["country"] + "|" + r["name_en"]  # noqa: E731

RANK_FILL = ["QS 2027 (место/диапазон)", "THE 2027 (место/диапазон)", "Webometrics (мир, место)",
             "Нац. рейтинг: название", "Нац. рейтинг: место", "Источник (URL)", "Комментарий"]
RANK_F = ["rank_qs", "rank_the", "rank_webometrics", None, None]
NC_FILL = ["Ректор (проверено)", "Почтовый адрес", "E-mail общий / ректората", "E-mail международного отдела",
           "Телефон", "ФИО адресата", "Должность адресата", "Обращение", "QS 2027", "THE 2027", "Webometrics (мир)",
           "Нац. рейтинг (название и место)", "Источник (URL)", "Комментарий"]


def build(outdir):
    rows = load()
    rows.sort(key=lambda r: (r["country"], r["name_en"]))
    conf = [r for r in rows if key(r) not in KEYS]
    nc = [r for r in rows if key(r) in KEYS]
    wb = Workbook(); wb.remove(wb.active)
    sheet(wb, "Рейтинги", ["Ключ", "Страна", "Город", "Вуз (EN)", "Вуз (родное)", "Сайт",
                           "QS сейчас", "THE сейчас", "Нац. сейчас"] + RANK_FILL,
          [4, 14, 16, 36, 28, 30, 12, 12, 14, 16, 16, 16, 22, 14, 40, 40],
          [[key(r), r["country"], r.get("city"), r["name_en"], r.get("name_native"), r.get("website"),
            r.get("rank_qs"), r.get("rank_the"), r.get("rank_national")] + [None] * 7 for r in conf],
          fill_cols=range(10, 17))
    wb.save(os.path.join(outdir, "BRICS_рейтинги_для_заполнения.xlsx"))
    wb = Workbook(); wb.remove(wb.active)
    sheet(wb, "Неподтверждённые", ["Ключ", "Страна", "Город", "Вуз (EN)", "Сайт", "Почему не подтверждён",
                                   "Ректор сейчас", "E-mail сейчас", "Волна"] + NC_FILL,
          [4, 14, 16, 36, 30, 40, 24, 28, 7] + [24, 34, 28, 28, 18, 24, 26, 22, 12, 12, 14, 24, 36, 40],
          [[key(r), r["country"], r.get("city"), r["name_en"], r.get("website"), KEYS[key(r)],
            r.get("head_name"), r.get("email_international") or r.get("email_general"), r["wave"]] + [None] * 14
           for r in nc], fill_cols=range(10, 24))
    wb.save(os.path.join(outdir, "BRICS_неподтверждённые_для_заполнения.xlsx"))
    print(len(conf), len(nc))


def add_src(r, src):
    cur = r.get("sources")
    if isinstance(cur, list):
        cur.append(src)
    else:
        r["sources"] = (cur + "\n" if cur else "") + src


def v(ws, row, col):
    x = ws.cell(row, col).value
    x = str(x).strip() if x is not None else ""
    return x or None


def import_filled(path):
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    mode = "rank" if ws.title == "Рейтинги" else "nc"
    db = {}
    for f in FILES:
        p = os.path.join(DATA, f + ".json")
        db[f] = json.load(open(p, encoding="utf-8"))
    idx = {key(r): r for rs in db.values() for r in rs}
    n = 0
    confirmed = []
    for i in range(2, ws.max_row + 1):
        r = idx.get(v(ws, i, 1))
        if not r:
            continue
        if mode == "rank":
            vals = [v(ws, i, c) for c in range(10, 17)]
            qs, the, web, nname, nplace, src, com = vals
            if not any(vals[:5]):
                if com:
                    r.setdefault("notes_ru", []).append("Рейтинги: " + com)
                continue
            if qs: r["rank_qs"] = "QS2027: " + qs
            if the: r["rank_the"] = "THE2027: " + the
            if web: r["rank_webometrics"] = "нет в опубл. списке" if web.startswith("нет") else web
            if nname and nname.startswith("нет"):
                r["rank_national"] = "нет общенационального рейтинга"
            elif nname and nplace and nplace != "—":
                r["rank_national"] = nname + ": " + nplace
            if src: add_src(r, src)
            if com: r.setdefault("notes_ru", []).append("Рейтинги: " + com)
        else:
            (head, addr, em, emi, tel, who, title, sal, qs, the, web, nat, src, com) = [v(ws, i, c) for c in range(10, 24)]
            if not any([head, addr, em, emi, tel, who, qs, the, web, nat]):
                if com:
                    r.setdefault("notes_ru", []).append("Ручная проверка: " + com)
                continue
            for k, x in (("head_name", head), ("postal_address", addr), ("email_general", em),
                         ("email_international", emi), ("phone", tel), ("addressee_name", who),
                         ("addressee_title", title), ("salutation", sal), ("rank_qs", qs), ("rank_the", the),
                         ("rank_webometrics", web), ("rank_national", nat)):
                if x: r[k] = x
            if emi or em:
                r["outreach_email_new"] = emi or em
            if src:
                add_src(r, src)
                r["head_verified_source"] = src
                r["confidence"] = "high"
                if head or emi or em:
                    confirmed.append(key(r))
            if com:
                r.setdefault("notes_ru", []).append("Ручная проверка: " + com)
        n += 1
    for f, rs in db.items():
        json.dump(rs, open(os.path.join(DATA, f + ".json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    if confirmed:
        for k in confirmed:
            KEYS.pop(k, None)
        json.dump(KEYS, open(os.path.join(DATA, "not_confirmed_keys.json"), "w", encoding="utf-8"),
                  ensure_ascii=False, indent=1)
    print("updated", n, "moved to confirmed", len(confirmed))


if __name__ == "__main__":
    {"build": lambda: build(sys.argv[2]), "import": lambda: import_filled(sys.argv[2])}[sys.argv[1]]()
