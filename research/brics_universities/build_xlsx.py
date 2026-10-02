"""Build the BRICS+ university outreach workbook from per-batch JSON research files.

Usage: python build_xlsx.py OUT.xlsx batch1.json [batch2.json ...]
"""
import datetime
import json
import re
import sys

import dns.resolver
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROFILE_RU = {
    "classical": "Классический",
    "technical/aerospace": "Технический / аэрокосмический",
    "agrarian/ecological": "Аграрный / экологический",
    "marine": "Морской / океанографический",
    "research institute": "НИИ / научный центр",
    "humanities": "Гуманитарный",
    "medical": "Медицинский",
    "other": "Другое",
}
COUNTRY_RU = {"China": "Китай"}
NEW_MEMBERS = {"Egypt", "Ethiopia", "Iran", "United Arab Emirates", "Indonesia", "Saudi Arabia"}
GROUND_SITES = {"South Africa", "Brazil"}

# (header, key or callable, width)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
# Mailboxes that exist but are poor targets for a partnership letter.
WEAK_MAILBOX = re.compile(r"^(admission|admissions|apply|zhaoban|webmaster|xwzx|news|studyat)", re.I)
RESTRICTED_RU = {
    "US Entity List": "Entity List США",
    "Seven Sons of National Defence": "«Семь сыновей национальной обороны»",
}


def bullets(text, sep=r"\s*;\s*"):
    """Render a delimited string (or list) as one bullet per line inside the cell."""
    if not text:
        return ""
    items = text if isinstance(text, list) else re.split(sep, text)
    items = [i.strip(" .") for i in items if i and i.strip(" .")]
    items = [i[0].upper() + i[1:] for i in items]
    return "\n".join(f"• {i}" for i in items) if len(items) > 1 else (items[0] if items else "")


def comma_bullets(text):
    # Split on commas/semicolons that are not inside parentheses.
    return bullets(text, r"\s*[;,]\s*(?![^()]*\))")


def outreach_email(r):
    """International office first (it handles foreign partnerships), then the general mailbox."""
    if r.get("outreach_email_new"):
        return EMAIL_RE.findall(r["outreach_email_new"])[0]
    found = EMAIL_RE.findall(r.get("email_international") or "") + EMAIL_RE.findall(r.get("email_general") or "")
    strong = [e for e in found if not WEAK_MAILBOX.match(e)]
    return (strong or found or [""])[0]


def addressee(r):
    """Name, title and office of whoever reads the outreach mailbox, plus the salutation to use."""
    if not outreach_email(r):
        return ""
    name = r.get("addressee_name") or "ФИО не найдено"
    lines = [name, r.get("addressee_title"), r.get("mailbox_owner"),
             f"Обращение: {r['salutation']}" if r.get("salutation") else ""]
    return "\n".join(l for l in lines if l)


def restricted_ru(text):
    for en, ru in RESTRICTED_RU.items():
        text = re.sub(re.escape(en), ru, text or "")
    return text.replace("added", "с").replace("; ", "\n")


# (header, key or callable, width)
COLUMNS = [
    ("№", None, 5),
    ("Страна", lambda r: COUNTRY_RU.get(r["country"], r["country"]), 10),
    ("Вуз (англ.)", "name_en", 34),
    ("Вуз (ориг.)", "name_native", 18),
    ("Профиль", lambda r: PROFILE_RU.get(r["profile_type"], r["profile_type"]), 20),
    ("Гос. / частный", "ownership", 24),
    ("Основные направления обучения", lambda r: comma_bullets(r.get("main_fields")), 40),
    ("Почтовый адрес", "postal_address", 40),
    ("E-mail для рассылки (рекомендуемый)", outreach_email, 28),
    ("Адресат (владелец рекомендуемого e-mail)", addressee, 36),
    ("E-mail (международный отдел)", "email_international", 26),
    ("E-mail (общий / ректорат)", "email_general", 26),
    ("Телефон", "phone", 18),
    ("Сайт", "website", 24),
    ("Ректор / президент", "head_name", 26),
    ("Должность", "head_title", 12),
    ("Секретарь парткома", "party_secretary", 24),
    ("Рейтинг QS", "rank_qs", 13),
    ("Рейтинг THE", "rank_the", 13),
    ("Национальный рейтинг", "rank_national", 16),
    ("Запущено КА (спутников)", "satellites_launched", 10),
    ("Названия запущенных КА (год)", lambda r: bullets(r.get("satellite_names")), 45),
    ("КА в разработке", lambda r: bullets(r.get("sats_in_development")), 30),
    ("Наземная станция", "ground_station", 30),
    ("Космические подразделения", lambda r: bullets(r.get("space_units")), 40),
    ("Партнёрства в космической сфере", lambda r: bullets(r.get("agency_partnerships")), 45),
    ("Санкционные / ограничительные списки (США)", lambda r: restricted_ru(r.get("restricted_list")), 22),
    ("Космический индекс (0–100)", "score", 11),
    ("Волна рассылки", "wave", 9),
    ("Проверка домена e-mail (MX)", "mx", 12),
    ("Достоверность", lambda r: {"high": "высокая", "medium": "средняя", "low": "низкая"}.get(r.get("confidence"), r.get("confidence")), 11),
    ("Примечания", lambda r: bullets(r.get("notes_ru") or r.get("notes")), 55),
    ("Источники", lambda r: bullets(r.get("sources")), 60),
    ("Дата проверки", "checked", 12),
]
WAVE_COL = next(i for i, (h, _, _) in enumerate(COLUMNS, 1) if h == "Волна рассылки")

NETWORKS = re.compile(r"UNISEC|IAF|International Astronautical Federation|BRICS|SCO|APSCO", re.I)


def best_rank(*texts):
    nums = []
    for t in texts:
        m = re.search(r":\s*=?(\d+)(?:\s*[-–]\s*(\d+))?", t or "")
        if m:
            nums.append(int(m.group(1)))
    return min(nums) if nums else None


def score(r):
    s = 0
    n = int(r.get("satellites_launched") or 0)
    if n:
        s += min(35 + 5 * (n - 1), 45)
    if r.get("sats_in_development"):
        s += 15
    if r.get("ground_station"):
        s += 10
    if r.get("space_units"):
        s += 10
    if r.get("agency_partnerships"):
        s += 10
    g = best_rank(r.get("rank_qs"), r.get("rank_the"))
    if g is not None:
        s += 10 if g <= 500 else 6 if g <= 1000 else 3
    nat = best_rank(r.get("rank_national"))
    if nat is not None and nat <= 60:
        s += 5
    if NETWORKS.search(r.get("agency_partnerships") or ""):
        s += 5
    return min(s, 100)


def wave(r):
    if int(r.get("satellites_launched") or 0) and r["score"] >= 60:
        return "1A"
    if r["country"] in NEW_MEMBERS | GROUND_SITES and r["score"] >= 40:
        return "1B"
    if r["score"] >= 35 or (r["profile_type"] == "technical/aerospace" and r.get("space_units")):
        return "2"
    return "3"


_mx_cache = {}


def mx_ok(*emails):
    """'да' if every published e-mail's domain has an MX record, 'нет' if one fails, '' if no e-mail."""
    doms = {d.lower() for e in emails if e for d in re.findall(r"[\w.+-]+@([\w-]+(?:\.[\w-]+)+)", e)}
    if not doms:
        return ""
    for d in doms:
        if d not in _mx_cache:
            try:
                dns.resolver.resolve(d, "MX", lifetime=4)
                _mx_cache[d] = True
            except Exception:
                try:
                    dns.resolver.resolve(d, "A", lifetime=4)  # implicit MX fallback
                    _mx_cache[d] = True
                except Exception:
                    _mx_cache[d] = False
        if not _mx_cache[d]:
            return "нет"
    return "да"


def write_sheet(ws, rows):
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    wave_fill = {"1A": "C6EFCE", "1B": "DDEBF7", "2": "FFF2CC", "3": "F2F2F2"}
    for c, (h, _, w) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    for i, r in enumerate(rows, 1):
        for c, (_, key, _) in enumerate(COLUMNS, 1):
            v = i if key is None else key(r) if callable(key) else r.get(key, "")
            cell = ws.cell(i + 1, c, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i + 1, WAVE_COL).fill = PatternFill("solid", fgColor=wave_fill[r["wave"]])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions


def main(out, paths):
    rows = []
    for p in paths:
        rows += json.load(open(p, encoding="utf-8"))
    today = datetime.date.today().isoformat()
    for r in rows:
        r["satellites_launched"] = int(r.get("satellites_launched") or 0)
        r["score"] = score(r)
        r["wave"] = wave(r)
        r["mx"] = mx_ok(r.get("email_general"), r.get("email_international"))
        r["checked"] = today
    rows.sort(key=lambda r: (r["country"], -r["score"], r["name_en"]))

    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    summary.append(["Университетская спутниковая группировка БРИКС: база вузов"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Дата формирования: {today}"])
    summary.append([])
    summary.append(["Страна", "Вузов", "Волна 1A", "Волна 1B", "Волна 2", "Волна 3", "С e-mail", "С ректором"])
    for c in summary[4]:
        c.font = Font(bold=True)
    for country in sorted({r["country"] for r in rows}):
        rs = [r for r in rows if r["country"] == country]
        summary.append([
            COUNTRY_RU.get(country, country), len(rs),
            *(sum(r["wave"] == w for r in rs) for w in ("1A", "1B", "2", "3")),
            sum(bool(r.get("email_general") or r.get("email_international")) for r in rs),
            sum(bool(r.get("head_name")) for r in rs),
        ])
    summary.append([])
    for line in [
        "Космический индекс (0–100): запущенные КА 35 (+5 за каждый следующий, макс. 45); КА в разработке 15;",
        "наземная станция 10; космические кафедры/институты 10; партнёрство с космическим агентством/отраслью 10;",
        "мировой рейтинг QS/THE (топ-500: 10, топ-1000: 6, в рейтинге: 3); национальный топ-60: 5; сети UNISEC/IAF/BRICS/SCO/APSCO: 5.",
        "Волна 1A: есть запущенные КА и индекс ≥ 60. Волна 1B: новые члены БРИКС+ и ключевые точки для наземных станций (индекс ≥ 40).",
        "Волна 2: индекс ≥ 35 или технический вуз с космическими подразделениями. Волна 3: потребители данных и образовательные партнёры.",
        "Пустая ячейка = данные не найдены в открытых официальных источниках (значения не домысливались).",
        "Проверка домена e-mail (MX): домен адреса принимает почту (да/нет).",
        "E-mail для рассылки: адрес международного отдела, при его отсутствии — общий; адреса приёмных комиссий и вебмастеров — только если других нет.",
        "Названия запущенных КА: спутники, созданные вузом (самостоятельно или совместно) и выведенные на орбиту, с годом запуска.",
        "Санкционные списки: Entity List Минторга США (экспортный контроль) и «Семь сыновей национальной обороны» (7 вузов КНР при MIIT, связанных с ОПК).",
        "Для партнёров из Индии, Бразилии, ОАЭ, ЮАР это может быть препятствием к участию в одном консорциуме.",
    ]:
        summary.append([line])
    summary.column_dimensions["A"].width = 18

    write_sheet(wb.create_sheet("Волна 1"), [r for r in rows if r["wave"] in ("1A", "1B")])
    for country in sorted({r["country"] for r in rows}):
        write_sheet(wb.create_sheet(COUNTRY_RU.get(country, country)), [r for r in rows if r["country"] == country])
    wb.save(out)

    csv_path = out.rsplit(".", 1)[0] + ".csv"
    import csv
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([h for h, _, _ in COLUMNS])
        for i, r in enumerate(rows, 1):
            w.writerow([i if k is None else k(r) if callable(k) else r.get(k, "") for _, k, _ in COLUMNS])
    print(f"{len(rows)} rows -> {out}, {csv_path}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2:])
