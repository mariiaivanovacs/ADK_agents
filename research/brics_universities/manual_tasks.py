"""Build a worklist for manual checking and import the filled worklist back into the data files.

  python manual_tasks.py build OUT.xlsx
  python manual_tasks.py import FILLED.xlsx
"""
import json
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
from build_xlsx import COUNTRY_RU, outreach_email, score, wave  # noqa: E402

DATA = os.path.join(os.path.dirname(__file__), "data")
FILES = ["china", "india", "brazil", "south_africa", "uae", "egypt", "iran", "indonesia",
         "saudi_arabia", "ethiopia", "partners", "russia"]
BLOCK = re.compile(r"недоступ|503|403|412|заблок|сброс|Cloudflare|SSL|TLS|таймаут|DNS|не открыл|не загруж|"
                   r"anti-bot|антибот|блокир|Incapsula|firewall|файрвол", re.I)
BOT = re.compile(r"Cloudflare|антибот|anti-bot|403|412|Incapsula|firewall|файрвол|блокир", re.I)
FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="1F3864")


def load():
    rows = []
    for f in FILES:
        for i, r in enumerate(json.load(open(os.path.join(DATA, f + ".json"), encoding="utf-8"))):
            r["_file"], r["_idx"] = f, i
            r["score"] = score(r)
            r["wave"] = wave(r)
            rows.append(r)
    return rows


def reason(r):
    notes = [n for n in (r.get("notes_ru") or []) if BLOCK.search(n)]
    text = " ".join(notes)
    kind = "защита от ботов (Cloudflare / 403 / 412)" if BOT.search(text) else "сервер не отвечает из нашей сети (503 / сброс / SSL / DNS)"
    return kind, (notes[0] if notes else "")[:260]


def sheet(wb, title, headers, widths, data, fill_cols=()):
    ws = wb.create_sheet(title)
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7F6000") if c in fill_cols else HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    for r, row in enumerate(data, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c in fill_cols:
                cell.fill = FILL
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def build(out):
    rows = load()
    order = {"1A": 0, "1B": 1, "2": 2, "3": 3}
    rows.sort(key=lambda r: (order[r["wave"]], r["country"], r["name_en"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Инструкция"
    for line in [
        "Рабочий файл для ручной проверки. Жёлтые колонки — для заполнения.",
        "",
        "Лист «Адресаты»: вузы, где адрес для рассылки есть (или нужен), но нет имени человека, которому адресуется письмо.",
        "  Заполните ФИО, должность и обращение (на языке письма). Если e-mail другой — впишите его в колонку «E-mail».",
        "Лист «Имя под вопросом»: имя указано, но достоверность низкая или имя могло смениться — подтвердите или исправьте.",
        "Лист «Сайты не открылись»: сайты, которые не открылись из нашей среды. Откройте и впишите ректора, адрес, e-mail, телефон.",
        "",
        "Волна 1A/1B — приоритет. Ключ строки: «Страна» + «Вуз (англ.)» — не меняйте эти колонки.",
        "Пустые жёлтые ячейки игнорируются. Заполненный файл пришлите обратно — я загружу данные в базу и пересоберу таблицы.",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 130

    # 1. addressees missing
    miss = [r for r in rows if not (r.get("addressee_name") or "").strip()]
    data = []
    for r in miss:
        data.append([r["wave"], COUNTRY_RU.get(r["country"], r["country"]), r["name_en"], r.get("website", ""),
                     outreach_email(r) or "— нет —", r.get("mailbox_owner", ""), r.get("salutation", ""),
                     "", "", "", ""])
    sheet(wb, "Адресаты", ["Волна", "Страна", "Вуз (англ.)", "Сайт", "E-mail для рассылки", "Чей ящик",
                           "Обращение сейчас", "ФИО адресата", "Должность", "Обращение (для письма)", "E-mail (если другой)"],
          [7, 14, 36, 28, 30, 30, 32, 28, 28, 30, 28], data, fill_cols=(8, 9, 10, 11))

    # 2. named but doubtful
    doubt = [r for r in rows if (r.get("addressee_name") or "").strip()
             and (r.get("confidence") == "low" or any("устарев" in n or "не подтвержд" in n for n in (r.get("notes_ru") or [])
                                                      if "дресат" in n))]
    data = [[r["wave"], COUNTRY_RU.get(r["country"], r["country"]), r["name_en"], r.get("website", ""),
             r.get("addressee_name", ""), r.get("addressee_title", ""), r.get("salutation", ""),
             outreach_email(r), "", "", ""] for r in doubt]
    sheet(wb, "Имя под вопросом", ["Волна", "Страна", "Вуз (англ.)", "Сайт", "Адресат сейчас", "Должность сейчас",
                                   "Обращение сейчас", "E-mail", "ФИО (исправленное)", "Должность (исправленная)", "Обращение (исправленное)"],
          [7, 14, 36, 28, 28, 28, 30, 30, 28, 28, 30], data, fill_cols=(9, 10, 11))

    # 3. sites not opened
    blocked = [r for r in rows if any(BLOCK.search(n) for n in (r.get("notes_ru") or []))]
    data = []
    for r in blocked:
        kind, note = reason(r)
        data.append([r["wave"], COUNTRY_RU.get(r["country"], r["country"]), r["name_en"], r.get("website", ""), kind,
                     note, r.get("head_name", ""), outreach_email(r), "", "", "", ""])
    sheet(wb, "Сайты не открылись", ["Волна", "Страна", "Вуз (англ.)", "Сайт", "Причина", "Что именно не вышло",
                                     "Ректор сейчас", "E-mail сейчас", "Ректор (проверено)", "Почтовый адрес", "E-mail", "Телефон"],
          [7, 14, 36, 28, 28, 50, 26, 28, 28, 34, 28, 18], data, fill_cols=(9, 10, 11, 12))
    wb.save(out)
    print(f"Адресаты: {len(miss)}; под вопросом: {len(doubt)}; сайты не открылись: {len(blocked)} -> {out}")


def val(ws, row, col):
    v = ws.cell(row, col).value
    return str(v).strip() if v not in (None, "") else ""


def import_filled(path):
    ru = {v: k for k, v in COUNTRY_RU.items()}
    by = {}
    for f in FILES:
        d = json.load(open(os.path.join(DATA, f + ".json"), encoding="utf-8"))
        by[f] = d
    index = {(r["country"], r["name_en"]): r for d in by.values() for r in d}
    wb = load_workbook(path)
    n = 0

    def find(ws, row):
        return index.get((ru.get(val(ws, row, 2), val(ws, row, 2)), val(ws, row, 3)))

    for title, cols in (("Адресаты", (8, 9, 10, 11)), ("Имя под вопросом", (9, 10, 11, None))):
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        for row in range(2, ws.max_row + 1):
            r = find(ws, row)
            if not r:
                continue
            name, ttl, sal = (val(ws, row, c) for c in cols[:3])
            mail = val(ws, row, cols[3]) if cols[3] else ""
            if not (name or ttl or sal or mail):
                continue
            if name:
                r["addressee_name"] = name
            if ttl:
                r["addressee_title"] = ttl
            if sal:
                r["salutation"] = sal
            if mail:
                r["outreach_email_new"] = mail
            r["confidence"] = "high" if r.get("confidence") == "low" else r.get("confidence")
            r.setdefault("notes_ru", []).append("Адресат внесён вручную (проверено человеком)")
            n += 1
    if "Сайты не открылись" in wb.sheetnames:
        ws = wb["Сайты не открылись"]
        for row in range(2, ws.max_row + 1):
            r = find(ws, row)
            if not r:
                continue
            head, addr, mail, phone = (val(ws, row, c) for c in (9, 10, 11, 12))
            if not (head or addr or mail or phone):
                continue
            if head:
                r["head_name"] = head
            if addr:
                r["postal_address"] = addr
            if mail:
                r["outreach_email_new"] = mail
            if phone:
                r["phone"] = phone
            r["confidence"] = "high"
            r.setdefault("notes_ru", []).append("Данные внесены вручную после проверки сайта человеком")
            n += 1
    for f, d in by.items():
        json.dump(d, open(os.path.join(DATA, f + ".json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"Обновлено строк: {n}")


if __name__ == "__main__":
    {"build": lambda: build(sys.argv[2]), "import": lambda: import_filled(sys.argv[2])}[sys.argv[1]]()
